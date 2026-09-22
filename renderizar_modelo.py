# renderizar_modelo.py
"""
Renderiza a reclamação trabalhista a partir do MODELO_PRINCIPAL v17.

A diferença em relação a deixar o LLM redigir: aqui o texto jurídico sai **literal**
do modelo do escritório. O LLM só decide quais blocos ligam e preenche as variáveis —
não reescreve tese, não inventa percentual e não omite capítulo.

Uso:
    python renderizar_modelo.py dados.json peca.docx
"""

import json
import re
import sys

from docx import Document

from config import BASE_DIR
import os

MODELO_V17 = os.path.join(BASE_DIR, "dados", "modelo_v17.docx")

RE_TAG_BLOCO = re.compile(r"\{\{([#^/])([a-zA-Z_0-9]+)\}\}")
RE_VARIAVEL = re.compile(r"\{\{([A-Za-z_0-9]+)\}\}")


def _ativo(marcador: str, valor) -> bool:
    """`{{#flag}}` entra quando há valor; `{{^flag}}` entra justamente quando não há."""
    return bool(valor) if marcador == "#" else not bool(valor)


def _resolver_blocos(documento, dados: dict) -> list[str]:
    """Aplica os blocos condicionais varrendo o documento como um fluxo contínuo.

    A tag pode estar sozinha no parágrafo, aberta e fechada no mesmo parágrafo, ou
    aberta no fim de um parágrafo e fechada vários adiante — por isso a pilha
    atravessa os parágrafos em vez de ser resolvida dentro de cada um.
    """
    pilha: list[bool] = []
    ativos: list[str] = []
    remover = []

    for paragrafo in documento.paragraphs:
        original = paragrafo.text
        if not original.strip() and not pilha:
            continue

        pedacos, posicao = [], 0
        for tag in RE_TAG_BLOCO.finditer(original):
            if all(pilha):
                pedacos.append(original[posicao:tag.start()])
            posicao = tag.end()

            marcador, flag = tag.groups()
            if marcador == "/":
                if pilha:
                    pilha.pop()
                continue

            ligado = _ativo(marcador, dados.get(flag))
            if ligado and marcador == "#" and all(pilha):
                ativos.append(flag)
            pilha.append(ligado and all(pilha))

        if all(pilha):
            pedacos.append(original[posicao:])

        novo = "".join(pedacos)
        if novo == original:
            continue
        if novo.strip():
            _reescrever(paragrafo, novo)
        else:
            remover.append(paragrafo)

    for paragrafo in remover:
        elemento = paragrafo._element
        elemento.getparent().remove(elemento)

    return ativos


def _reescrever(paragrafo, texto: str):
    """Troca o conteúdo do parágrafo mantendo a formatação do primeiro run."""
    if paragrafo.runs:
        paragrafo.runs[0].text = texto
        for run in paragrafo.runs[1:]:
            run.text = ""
    else:
        paragrafo.add_run(texto)


def renderizar(dados: dict, destino: str, modelo: str = MODELO_V17) -> dict:
    """Preenche o modelo e grava o .docx. Devolve o relatório do que entrou e do que falta."""
    if not os.path.exists(modelo):
        raise SystemExit(f"❌ modelo não encontrado: {modelo}")

    documento = Document(modelo)
    blocos_ativos = _resolver_blocos(documento, dados)

    faltando, preenchidas = set(), set()

    for paragrafo in documento.paragraphs:
        if "{{" not in paragrafo.text:
            continue
        for run in paragrafo.runs:
            if "{{" not in run.text:
                continue

            def troca(achado):
                nome = achado.group(1)
                valor = dados.get(nome)
                if valor in (None, ""):
                    faltando.add(nome)
                    return f"[{nome} — PREENCHER]"
                preenchidas.add(nome)
                return str(valor)

            run.text = RE_VARIAVEL.sub(troca, run.text)

    documento.save(destino)
    return {
        "arquivo": destino,
        "blocos_ativos": sorted(set(blocos_ativos)),
        "variaveis_preenchidas": len(preenchidas),
        "variaveis_faltando": sorted(faltando),
    }


if __name__ == "__main__":
    if len(sys.argv) < 3:
        raise SystemExit("uso: python renderizar_modelo.py <dados.json> <saida.docx>")
    with open(sys.argv[1], "r", encoding="utf-8") as arquivo:
        relatorio = renderizar(json.load(arquivo), sys.argv[2])
    print(f"✅ {relatorio['arquivo']}")
    print(f"   blocos ativos: {', '.join(relatorio['blocos_ativos'])}")
    print(f"   variáveis preenchidas: {relatorio['variaveis_preenchidas']}")
    if relatorio["variaveis_faltando"]:
        print(f"   ⚠️  faltando: {', '.join(relatorio['variaveis_faltando'])}")
