# gerar_docx.py
"""
Converte a peça em Markdown para .docx sobre o papel timbrado do escritório.

O timbre (logo no cabeçalho e barra de contatos no rodapé) são imagens dentro do
`dados/modelo_peca.docx` — por isso a entrega final é .docx, e não Markdown: nenhum
texto puro carrega logo.

Uso:
    python gerar_docx.py peca.md peca.docx
"""

import os
import re
import sys

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

from config import BASE_DIR

MODELO = os.path.join(BASE_DIR, "dados", "modelo_peca.docx")

FONTE = "Arial"
TAMANHO = Pt(12)


def _escrever_runs(paragrafo, texto: str):
    """Escreve o texto aplicando negrito onde o Markdown usa **."""
    for i, parte in enumerate(re.split(r"\*\*(.+?)\*\*", texto)):
        if not parte:
            continue
        run = paragrafo.add_run(parte)
        run.bold = i % 2 == 1  # as partes ímpares vieram entre **
        run.font.name = FONTE
        run.font.size = TAMANHO


def markdown_para_docx(markdown: str, destino: str, modelo: str = MODELO) -> str:
    """Gera o .docx da peça preservando cabeçalho, rodapé e margens do timbre."""
    if not os.path.exists(modelo):
        raise SystemExit(f"❌ modelo não encontrado: {modelo}")

    documento = Document(modelo)
    markdown = re.sub(r"^```\w*\n|\n```$", "", markdown.strip())

    for linha in markdown.split("\n"):
        linha = linha.rstrip()

        if not linha:
            continue

        if re.match(r"^#{1,2}\s", linha):  # título da peça e capítulos
            paragrafo = documento.add_paragraph()
            paragrafo.alignment = (
                WD_ALIGN_PARAGRAPH.CENTER if linha.startswith("# ") else WD_ALIGN_PARAGRAPH.LEFT
            )
            paragrafo.paragraph_format.space_before = Pt(12)
            paragrafo.paragraph_format.space_after = Pt(6)
            _escrever_runs(paragrafo, f"**{re.sub(r'^#+ *', '', linha).upper()}**")
            continue

        if re.match(r"^\s*([-*]|\d+\.)\s", linha):  # itens de lista e pedidos
            paragrafo = documento.add_paragraph()
            paragrafo.paragraph_format.left_indent = Pt(24)
            paragrafo.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            _escrever_runs(paragrafo, re.sub(r"^\s*[-*]\s", "• ", linha))
            continue

        if re.match(r"^-{3,}$", linha):  # separador horizontal do Markdown
            continue

        paragrafo = documento.add_paragraph()
        paragrafo.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        paragrafo.paragraph_format.space_after = Pt(6)
        _escrever_runs(paragrafo, linha)

    documento.save(destino)
    return destino


if __name__ == "__main__":
    if len(sys.argv) < 3:
        raise SystemExit("uso: python gerar_docx.py <entrada.md> <saida.docx>")
    entrada, saida = sys.argv[1], sys.argv[2]
    with open(entrada, "r", encoding="utf-8") as arquivo:
        print(f"✅ {markdown_para_docx(arquivo.read(), saida)}")
